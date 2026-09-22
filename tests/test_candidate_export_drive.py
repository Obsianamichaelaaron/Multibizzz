import io
import openpyxl
from unittest.mock import patch, MagicMock
from django.test import TestCase, Client, override_settings
from django.urls import reverse
from django.utils import timezone

from accounts.models import User, ApplicantProfile, EmployerProfile
from jobs.models import JobPosting, EmployerRequest, CandidateExport
from applications.models import Application
from jobs.services.excel_exporter import generate_candidates_excel
from jobs.services.google_drive import upload_candidates_excel_to_drive, get_google_drive_service


@override_settings(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
class CandidateExportAndDriveWorkflowTest(TestCase):
    def setUp(self):
        self.client = Client()

        # 1. Users
        self.admin_user = User.objects.create_superuser(
            email='admin@multibiz.com',
            password='Password123!',
            first_name='Admin',
            last_name='User',
            role='admin'
        )

        self.employer_user_a = User.objects.create_user(
            email='employerA@techcorp.com',
            password='Password123!',
            first_name='Alice',
            last_name='Boss',
            role='employer'
        )
        self.employer_profile_a = EmployerProfile.objects.create(
            user=self.employer_user_a,
            company_name='TechCorp Solutions',
            industry='Technology'
        )

        self.employer_user_b = User.objects.create_user(
            email='employerB@othercorp.com',
            password='Password123!',
            first_name='Bob',
            last_name='Other',
            role='employer'
        )
        self.employer_profile_b = EmployerProfile.objects.create(
            user=self.employer_user_b,
            company_name='OtherCorp Inc',
            industry='Finance'
        )

        # 2. Talent Request & Job for Employer A
        self.talent_request_a = EmployerRequest.objects.create(
            employer=self.employer_profile_a,
            title='Senior Python Engineer',
            description='Build scalable Django applications and APIs',
            skills_required='Python, Django, PostgreSQL, REST API',
            experience_required='3-5 years',
            education_level='Bachelor\'s Degree',
            vacancies_count=2,
            status='submitted'
        )

        self.job_a = JobPosting.objects.create(
            employer=self.employer_profile_a,
            talent_request=self.talent_request_a,
            title='Senior Python Engineer',
            description='Build scalable Django applications and APIs',
            location='Manila, Philippines',
            employment_type='full-time',
            skills_required='Python, Django, PostgreSQL, REST API',
            min_experience_years=3,
            required_education='Bachelor\'s Degree',
            status='active'
        )

        # 3. Candidates: 1 Qualified, 1 Under-Qualified, 1 Not-Qualified
        # Candidate 1: Qualified
        self.user_cand_1 = User.objects.create_user(
            email='cand1@example.com',
            password='Password123!',
            first_name='John',
            last_name='Pythonista',
            role='applicant',
            phone='+639171112222'
        )
        self.profile_cand_1 = ApplicantProfile.objects.create(
            user=self.user_cand_1,
            skills='Python, Django, PostgreSQL, REST API, Docker',
            experience_years=4,
            education_level='Bachelor\'s Degree'
        )
        self.app_cand_1 = Application.objects.create(
            job=self.job_a,
            applicant=self.profile_cand_1,
            match_score=92.0,
            skills_match_score=95.0,
            exp_match_score=90.0,
            edu_match_score=90.0,
            qualification_status='qualified',
            qualification_reason='Strong direct Python & Django background.',
            status='pending'
        )

        # Candidate 2: Under-Qualified
        self.user_cand_2 = User.objects.create_user(
            email='cand2@example.com',
            password='Password123!',
            first_name='Jane',
            last_name='Junior',
            role='applicant',
            phone='+639173334444'
        )
        self.profile_cand_2 = ApplicantProfile.objects.create(
            user=self.user_cand_2,
            skills='Python, Flask, SQLite',
            experience_years=1,
            education_level='Associate Degree'
        )
        self.app_cand_2 = Application.objects.create(
            job=self.job_a,
            applicant=self.profile_cand_2,
            match_score=55.0,
            skills_match_score=60.0,
            exp_match_score=40.0,
            edu_match_score=65.0,
            qualification_status='under_qualified',
            qualification_reason='Transferable Python experience, entry level.',
            status='pending'
        )

        # Candidate 3: Not-Qualified
        self.user_cand_3 = User.objects.create_user(
            email='cand3@example.com',
            password='Password123!',
            first_name='Sam',
            last_name='Sales',
            role='applicant',
            phone='+639175556666'
        )
        self.profile_cand_3 = ApplicantProfile.objects.create(
            user=self.user_cand_3,
            skills='Sales, Cold Calling, CRM',
            experience_years=5,
            education_level='Bachelor\'s Degree'
        )
        self.app_cand_3 = Application.objects.create(
            job=self.job_a,
            applicant=self.profile_cand_3,
            match_score=20.0,
            skills_match_score=10.0,
            exp_match_score=30.0,
            edu_match_score=20.0,
            qualification_status='not_qualified',
            qualification_reason='No relevant software engineering skills.',
            status='pending'
        )

    def test_excel_exporter_generates_valid_workbook_with_all_candidates(self):
        """Verify that Excel exporter includes both Qualified and Not Qualified candidates with full metadata."""
        apps = Application.objects.filter(job=self.job_a).order_by('-match_score')
        excel_buffer, filename = generate_candidates_excel(self.job_a, apps)

        self.assertTrue(filename.endswith('.xlsx'))
        self.assertIn('Python', filename)

        # Load workbook from buffer and inspect
        wb = openpyxl.load_workbook(excel_buffer)
        ws = wb.active
        self.assertEqual(ws.title, "Candidate Evaluations")

        # Header banner checks
        self.assertIn("MULTIBIZ CANDIDATE EVALUATION REPORT", ws["A1"].value)
        self.assertIn("TechCorp Solutions", ws["A2"].value)
        self.assertIn("Total Candidates: 3", ws["A2"].value)

        # Row 4: Column Headers
        headers = [ws.cell(row=4, column=c).value for c in range(1, 14)]
        self.assertIn("Candidate Name", headers)
        self.assertIn("Qualification Status", headers)
        self.assertIn("Match Score", headers)
        self.assertIn("Matched Skills", headers)

        # Rows 5 to 7: Check candidate names and statuses
        row_names = [ws.cell(row=r, column=2).value for r in range(5, 8)]
        self.assertIn("John Pythonista", row_names)
        self.assertIn("Jane Junior", row_names)
        self.assertIn("Sam Sales", row_names)

        statuses = [ws.cell(row=r, column=6).value for r in range(5, 8)]
        self.assertIn("QUALIFIED", statuses)
        self.assertIn("UNDER QUALIFIED", statuses)
        self.assertIn("NOT QUALIFIED", statuses)

    def test_upload_candidates_excel_to_drive_sets_reader_permission(self):
        """Test Google Drive upload service with mocked Google API client to ensure role='reader' (view-only)."""
        mock_drive_service = MagicMock()
        mock_files = MagicMock()
        mock_permissions = MagicMock()

        mock_drive_service.files.return_value = mock_files
        mock_drive_service.permissions.return_value = mock_permissions

        mock_files.create.return_value.execute.return_value = {
            'id': 'google_drive_file_12345',
            'webViewLink': 'https://drive.google.com/file/d/google_drive_file_12345/view?usp=sharing'
        }
        mock_permissions.create.return_value.execute.return_value = {'id': 'perm_12345'}

        fake_excel_bytes = b"Fake Excel Content"

        with patch('jobs.services.google_drive.get_google_drive_service', return_value=mock_drive_service):
            result = upload_candidates_excel_to_drive(
                file_content_bytes=fake_excel_bytes,
                filename="Candidates_Senior_Python_Engineer.xlsx",
                employer_email="employerA@techcorp.com"
            )

            self.assertTrue(result['success'])
            self.assertEqual(result['file_id'], 'google_drive_file_12345')
            self.assertIn('google_drive_file_12345', result['web_view_link'])

            # Verify permissions.create was called with STRICT role='reader' (View-Only)
            mock_permissions.create.assert_called_once()
            call_kwargs = mock_permissions.create.call_args[1]
            self.assertEqual(call_kwargs['fileId'], 'google_drive_file_12345')
            self.assertEqual(call_kwargs['body']['type'], 'user')
            self.assertEqual(call_kwargs['body']['role'], 'reader')
            self.assertEqual(call_kwargs['body']['emailAddress'], 'employerA@techcorp.com')

    def test_upload_candidates_excel_sandbox_fallback(self):
        """Test sandbox fallback when Google API is not configured."""
        fake_excel_bytes = b"Fake Excel Content Sandbox"
        with patch('jobs.services.google_drive.get_google_drive_service', return_value=None):
            result = upload_candidates_excel_to_drive(
                file_content_bytes=fake_excel_bytes,
                filename="Candidates_Test.xlsx",
                employer_email="employerA@techcorp.com"
            )
            self.assertTrue(result['success'])
            self.assertTrue(result['is_mock'])
            self.assertEqual(result['web_view_link'], '')

    @patch('applications.views.send_candidates_sent_to_employer_email')
    def test_admin_send_candidates_to_employer_full_workflow(self, mock_email):
        """
        Verify the Admin dispatch & export workflow:
        1. Generates Excel including Qualified & Not Qualified.
        2. Uploads to Drive with view-only.
        3. Creates CandidateExport record.
        4. Updates status to sent_to_employer.
        5. Displays the required success message.
        """
        mock_email.return_value = True
        self.client.force_login(self.admin_user)
        url = reverse('admin_send_candidates_to_employer', kwargs={'pk': self.job_a.id})

        response = self.client.post(url, follow=True)

        self.assertEqual(response.status_code, 200)

        # 1. Required exact message confirmation
        expected_msg = "Candidates successfully exported and sent to the employer. The Excel file has been saved to Google Drive with view-only access."
        messages_list = [m.message for m in response.context['messages']]
        self.assertTrue(any(expected_msg in m for m in messages_list))

        # 2. Database CandidateExport record
        export_record = CandidateExport.objects.filter(job=self.job_a).first()
        self.assertIsNotNone(export_record)
        self.assertEqual(export_record.employer, self.employer_profile_a)
        self.assertEqual(export_record.total_candidates, 3)
        self.assertEqual(export_record.qualified_count, 1)
        self.assertEqual(export_record.not_qualified_count, 2)
        self.assertEqual(export_record.google_drive_view_link, '')
        self.assertTrue(bool(export_record.excel_file))

        # 3. Qualified candidates updated
        self.app_cand_1.refresh_from_db()
        self.assertTrue(self.app_cand_1.sent_to_employer)
        self.assertEqual(self.app_cand_1.radar_rank, 1)

        # 4. Talent request updated
        self.talent_request_a.refresh_from_db()
        self.assertEqual(self.talent_request_a.status, 'candidates_sent')

    def test_admin_export_candidates_excel_direct_download(self):
        """Test admin direct download of .xlsx candidate report."""
        self.client.force_login(self.admin_user)
        url = reverse('admin_export_candidates_excel', kwargs={'pk': self.job_a.id})

        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('attachment; filename="Candidates_', response['Content-Disposition'])

        # Verify content is a valid xlsx
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.active.title, "Candidate Evaluations")

    def test_employer_data_isolation_on_exports(self):
        """
        Verify strict Employer Data Isolation:
        Employer A can access Employer A's export.
        Employer B CANNOT access Employer A's export (403 Forbidden).
        """
        export_a = CandidateExport.objects.create(
            job=self.job_a,
            employer=self.employer_profile_a,
            talent_request=self.talent_request_a,
            filename='Candidates_Export_A.xlsx',
            google_drive_file_id='drive_file_a_123',
            google_drive_view_link='https://drive.google.com/file/d/drive_file_a_123/view?usp=sharing',
            exported_by=self.admin_user,
            total_candidates=3,
            qualified_count=1,
            not_qualified_count=2,
            sent_to_employer=True
        )

        view_url = reverse('employer_view_candidate_export', kwargs={'export_id': export_a.id})

        # 1. Employer A (Owner) accesses export -> Redirects to Drive View Link
        self.client.force_login(self.employer_user_a)
        resp_a = self.client.get(view_url)
        self.assertEqual(resp_a.status_code, 302)
        self.assertEqual(resp_a['Location'], export_a.google_drive_view_link)

        # 2. Employer B (Unauthorized) tries to access export -> 403 Forbidden
        self.client.force_login(self.employer_user_b)
        resp_b = self.client.get(view_url)
        self.assertEqual(resp_b.status_code, 403)

        # 3. Applicant tries to access export -> 403 Forbidden
        self.client.force_login(self.user_cand_1)
        resp_c = self.client.get(view_url)
        self.assertEqual(resp_c.status_code, 403)

        # 4. Admin accesses export -> Allowed (redirects to Drive link)
        self.client.force_login(self.admin_user)
        resp_admin = self.client.get(view_url)
        self.assertEqual(resp_admin.status_code, 302)
        self.assertEqual(resp_admin['Location'], export_a.google_drive_view_link)
